import streamlit as st
from io import BytesIO
import html
import re
import base64
from datetime import date

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.colors import HexColor
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image as RLImage,
    HRFlowable,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


# =========================================================
# CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="CV Pro Studio",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="auto",
)


# =========================================================
# TEXTES
# =========================================================

TEXTES = {
    "Français": {
        "nom": "Nom complet",
        "titre": "Titre professionnel",
        "email": "E-mail",
        "telephone": "Téléphone",
        "adresse": "Adresse",
        "linkedin": "LinkedIn",
        "site": "Site web",
        "profil": "Profil professionnel",
        "formation": "Formations",
        "experience": "Expériences professionnelles",
        "competence": "Compétences",
        "langue": "Langues",
        "certification": "Certifications",
        "projet": "Projets",
        "association": "Associations et engagements",
        "interet": "Intérêts et loisirs",
    },
    "English": {
        "nom": "Full name",
        "titre": "Professional title",
        "email": "Email",
        "telephone": "Phone",
        "adresse": "Address",
        "linkedin": "LinkedIn",
        "site": "Website",
        "profil": "Professional profile",
        "formation": "Education",
        "experience": "Professional experience",
        "competence": "Skills",
        "langue": "Languages",
        "certification": "Certifications",
        "projet": "Projects",
        "association": "Associations and volunteering",
        "interet": "Interests and hobbies",
    },
}


# =========================================================
# MODÈLES ET COULEURS
# =========================================================

MODELES = [
    "ATS",
    "Professionnel",
    "Européen",
    "International",
    "Académique",
    "Moderne",
    "Créatif",
    "Élégant",
    "Corporate",
    "Minimal",
]


COULEURS = {
    "Bleu professionnel": "#1F4E79",
    "Bleu moderne": "#2563EB",
    "Vert": "#16805B",
    "Rouge": "#B42318",
    "Violet": "#6D4AFF",
    "Orange": "#C65D00",
    "Gris anthracite": "#374151",
    "Noir": "#111111",
}


NIVEAUX = {
    "Français": [
        "Débutant",
        "Intermédiaire",
        "Avancé",
        "Expert",
    ],
    "English": [
        "Beginner",
        "Intermediate",
        "Advanced",
        "Expert",
    ],
}


MOIS_FR = [
    "Janvier",
    "Février",
    "Mars",
    "Avril",
    "Mai",
    "Juin",
    "Juillet",
    "Août",
    "Septembre",
    "Octobre",
    "Novembre",
    "Décembre",
]


# =========================================================
# SESSION STATE
# =========================================================

DEFAULTS = {
    "nom": "",
    "titre": "",
    "email": "",
    "telephone": "",
    "adresse": "",
    "linkedin": "",
    "site": "",
    "profil": "",
    "photo": None,
    "formation": [],
    "experiences": [],
    "competences": [],
    "langues": [],
    "certifications": [],
    "projets": [],
    "associations": [],
    "interets": [],
    "offre": "",
    "modele": "Professionnel",
    "couleur": "#1F4E79",
    "langue_cv": "Français",
}


for key, value in DEFAULTS.items():
    if key not in st.session_state:
        st.session_state[key] = value


# =========================================================
# FONCTIONS UTILITAIRES
# =========================================================

def esc(value):
    """Échappe le HTML."""
    if value is None:
        return ""
    return html.escape(str(value))


def couleur_rgb(hex_color):
    """Convertit une couleur HEX en RGBColor pour Word."""
    hex_color = hex_color.replace("#", "")

    return RGBColor(
        int(hex_color[0:2], 16),
        int(hex_color[2:4], 16),
        int(hex_color[4:6], 16),
    )


def format_mois_annee(mois, annee):
    return f"{mois} {annee}"


def choisir_mois_annee(label, key_prefix):
    """Sélection du mois et de l'année."""

    col1, col2 = st.columns(2)

    annee_actuelle = date.today().year

    with col1:
        mois = st.selectbox(
            f"{label} – Mois",
            MOIS_FR,
            index=0,
            key=f"{key_prefix}_mois",
        )

    with col2:
        annees = list(range(annee_actuelle + 10, 1949, -1))

        annee = st.selectbox(
            f"{label} – Année",
            annees,
            index=10,
            key=f"{key_prefix}_annee",
        )

    return mois, annee


def section_card(numero, titre, couleur):
    """Affiche une carte de section."""

    st.html(
        f"""
        <div class="section-card"
             style="
                border-left: 6px solid {couleur};
                border-top: 1px solid {couleur}30;
                border-right: 1px solid {couleur}30;
                border-bottom: 1px solid {couleur}30;
             ">

            <div class="section-number"
                 style="background:{couleur};">
                {numero}
            </div>

            <div>
                <div class="section-title">
                    {esc(titre)}
                </div>

                <div class="section-line"
                     style="background:{couleur};">
                </div>
            </div>
        </div>
        """
    )


# =========================================================
# COULEUR ACTUELLE
# =========================================================

COULEUR = st.session_state.couleur


# =========================================================
# CSS
# =========================================================

st.markdown(
    f"""
    <style>

    /* =====================================================
       APPLICATION
       ===================================================== */

    .stApp {{
        background: #f5f7fb;
    }}

    .main .block-container {{
        max-width: 1450px;
        padding-top: 2rem;
        padding-bottom: 4rem;
    }}


    /* =====================================================
       HEADER
       ===================================================== */

    .app-header {{
        background: linear-gradient(
            135deg,
            {COULEUR},
            #111827
        );

        padding: 30px;
        border-radius: 20px;
        margin-bottom: 25px;

        box-shadow:
            0 10px 30px rgba(0,0,0,0.12);

        color: white;
    }}

    .app-header h1 {{
        margin: 0;
        font-size: 34px;
        font-weight: 800;
    }}

    .app-header p {{
        margin-top: 8px;
        margin-bottom: 0;
        opacity: 0.9;
        font-size: 16px;
    }}


    /* =====================================================
       CARTES DE SECTION
       ===================================================== */

    .section-card {{
        background: white;

        border-radius: 15px;

        padding: 18px 20px;

        margin-top: 10px;
        margin-bottom: 22px;

        display: flex;

        align-items: center;

        gap: 15px;

        box-shadow:
            0 5px 18px rgba(0,0,0,0.06);
    }}

    .section-number {{
        width: 44px;
        height: 44px;

        min-width: 44px;

        border-radius: 12px;

        display: flex;

        align-items: center;
        justify-content: center;

        color: white;

        font-size: 18px;
        font-weight: 800;
    }}

    .section-title {{
        font-size: 22px;

        font-weight: 800;

        color: #111827;
    }}

    .section-line {{
        width: 50px;
        height: 3px;

        border-radius: 10px;

        margin-top: 5px;
    }}


    /* =====================================================
       CARTES INFORMATIONS
       ===================================================== */

    .info-card {{
        background: white;

        padding: 20px;

        border-radius: 15px;

        border: 1px solid #e5e7eb;

        box-shadow:
            0 4px 15px rgba(0,0,0,0.04);
    }}


    /* =====================================================
       SIDEBAR
       ===================================================== */

    section[data-testid="stSidebar"] {{
        background: #111827;
    }}

    section[data-testid="stSidebar"] * {{
        color: white;
    }}


    /* =====================================================
       INPUTS
       ===================================================== */

    div[data-baseweb="input"],
    div[data-baseweb="textarea"],
    div[data-baseweb="select"] {{
        border-radius: 10px;
    }}


    /* =====================================================
       BOUTONS
       ===================================================== */

    .stButton > button,
    .stDownloadButton > button {{
        border-radius: 10px;

        font-weight: 700;

        min-height: 42px;
    }}


    /* =====================================================
       CV
       ===================================================== */

    .cv-paper {{
        background: white;

        width: 794px;

        max-width: 100%;

        min-height: 1123px;

        box-sizing: border-box;

        padding: 48px 52px;

        border-radius: 4px;

        box-shadow:
            0 14px 40px rgba(0,0,0,0.14);

        margin: 20px auto 35px auto;

        color: #111827;

        font-family:
            Arial,
            Helvetica,
            sans-serif;

        line-height: 1.5;

        overflow-wrap: anywhere;
    }}

    .cv-name {{
        font-size: 31px;

        font-weight: 800;

        margin-bottom: 4px;
    }}

    .cv-title {{
        font-size: 17px;

        margin-bottom: 15px;
    }}

    .cv-section-title {{
        font-size: 15px;

        font-weight: 800;

        margin-top: 24px;

        margin-bottom: 5px;

        text-transform: uppercase;

        letter-spacing: 0.7px;
    }}

    .cv-line {{
        height: 2px;

        width: 100%;

        margin-bottom: 11px;

        border-radius: 2px;
    }}

    .cv-item {{
        margin-bottom: 13px;
    }}

    .cv-date {{
        font-weight: 700;

        color: #4b5563;

        font-size: 13px;
    }}

    .cv-small {{
        color: #4b5563;

        font-size: 13px;
    }}

    .cv-contact {{
        color: #374151;

        font-size: 13px;

        line-height: 1.8;
    }}

    .skill {{
        display: inline-block;

        padding: 5px 10px;

        margin: 3px;

        border-radius: 15px;

        background: #f1f5f9;

        font-size: 12px;

        font-weight: 600;
    }}


    /* =====================================================
       MOBILE
       ===================================================== */

    @media (max-width: 850px) {{

        .main .block-container {{
            padding:
                1rem
                0.75rem
                3rem
                0.75rem;

            max-width: 100%;
        }}

        .app-header {{
            padding: 22px 18px;

            border-radius: 14px;

            margin-bottom: 16px;
        }}

        .app-header h1 {{
            font-size: 25px;

            line-height: 1.2;
        }}

        .app-header p {{
            font-size: 14px;

            line-height: 1.4;
        }}

        .section-card {{
            padding: 13px 14px;

            margin-top: 6px;

            margin-bottom: 16px;

            border-radius: 12px;

            gap: 10px;
        }}

        .section-number {{
            width: 36px;

            height: 36px;

            min-width: 36px;

            border-radius: 9px;

            font-size: 15px;
        }}

        .section-title {{
            font-size: 17px;

            line-height: 1.25;
        }}

        div[data-testid="stHorizontalBlock"] {{
            gap: 0.7rem;
        }}

        input,
        textarea {{
            font-size: 16px !important;
        }}

        .stButton > button,
        .stDownloadButton > button {{
            min-height: 48px;

            font-size: 15px;

            border-radius: 10px;
        }}

        div[data-testid="stExpander"] {{
            border-radius: 12px;

            margin-bottom: 10px;
        }}

        .cv-paper {{
            width: 100%;

            max-width: 100%;

            min-height: auto;

            padding: 25px 20px;

            margin: 12px auto 25px auto;

            border-radius: 3px;

            box-shadow:
                0 8px 25px rgba(0,0,0,0.12);

            overflow-wrap: break-word;

            word-break: normal;
        }}

        .cv-name {{
            font-size: 25px;

            line-height: 1.15;
        }}

        .cv-title {{
            font-size: 14px;

            line-height: 1.4;
        }}

        .cv-section-title {{
            font-size: 13px;

            margin-top: 19px;

            line-height: 1.3;
        }}

        .cv-contact {{
            font-size: 11px;

            line-height: 1.7;

            word-break: break-word;
        }}

        .cv-date {{
            font-size: 11px;
        }}

        .cv-small {{
            font-size: 11px;

            line-height: 1.5;
        }}

        .cv-item {{
            margin-bottom: 11px;
        }}

        .skill {{
            font-size: 10px;

            padding: 5px 8px;

            margin: 2px;
        }}

        .cv-paper img {{
            width: 80px !important;

            height: 80px !important;

            margin-left: 12px !important;
        }}

        button[data-baseweb="tab"] {{
            font-size: 12px !important;

            padding-left: 8px !important;

            padding-right: 8px !important;
        }}

        div[data-testid="stTabs"] > div:first-child {{
            overflow-x: auto;

            scrollbar-width: thin;
        }}

        div[data-testid="stMetric"] {{
            padding: 10px;
        }}
    }}

    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# SIDEBAR
# =========================================================

with st.sidebar:

    st.markdown("## 📄 CV PRO STUDIO")

    st.caption(
        "Créateur de CV professionnel"
    )

    st.markdown("---")

    modele = st.selectbox(
        "🎨 Modèle de CV",
        MODELES,
        index=MODELES.index(
            st.session_state.modele
        ),
    )

    st.session_state.modele = modele

    couleur_nom = st.selectbox(
        "🎨 Couleur principale",
        list(COULEURS.keys()),
        index=list(
            COULEURS.values()
        ).index(
            st.session_state.couleur
        ),
    )

    st.session_state.couleur = COULEURS[
        couleur_nom
    ]

    langue_cv = st.selectbox(
        "🌐 Langue du CV",
        ["Français", "English"],
        index=(
            0
            if st.session_state.langue_cv
            == "Français"
            else 1
        ),
    )

    st.session_state.langue_cv = langue_cv

    st.markdown("---")

    st.markdown("### Sections")

    st.markdown(
        """
        ① Informations  
        ② Profil  
        ③ Formations  
        ④ Expériences  
        ⑤ Compétences  
        ⑥ Langues  
        ⑦ Certifications  
        ⑧ Projets  
        ⑨ Associations  
        ⑩ Intérêts  
        ⑪ Offre  
        ⑫ Aperçu & Export
        """
    )


# =========================================================
# HEADER
# =========================================================

st.html(
    """
    <div class="app-header">

        <h1>
            CV Pro Studio
        </h1>

        <p>
            Créez un CV professionnel,
            moderne et prêt à envoyer.
        </p>

    </div>
    """
)


# =========================================================
# ONGLETS
# =========================================================

tabs = st.tabs(
    [
        "① Informations",
        "② Profil",
        "③ Formations",
        "④ Expériences",
        "⑤ Compétences",
        "⑥ Langues",
        "⑦ Certifications",
        "⑧ Projets",
        "⑨ Associations",
        "⑩ Intérêts",
        "⑪ Offre",
        "⑫ Aperçu & Export",
    ]
)


# =========================================================
# 1 - INFORMATIONS
# =========================================================

with tabs[0]:

    section_card(
        "1",
        "Informations personnelles",
        COULEUR,
    )

    col1, col2 = st.columns(
        [2, 1]
    )

    with col1:

        st.session_state.nom = st.text_input(
            "Nom complet",
            value=st.session_state.nom,
            placeholder="Exemple : Amani Francis Kanga",
        )

        st.session_state.titre = st.text_input(
            "Titre professionnel",
            value=st.session_state.titre,
            placeholder="Exemple : Analyste en politiques environnementales",
        )

        st.session_state.email = st.text_input(
            "E-mail",
            value=st.session_state.email,
            placeholder="exemple@email.com",
        )

        st.session_state.telephone = st.text_input(
            "Téléphone",
            value=st.session_state.telephone,
            placeholder="+225 XX XX XX XX XX",
        )

        st.session_state.adresse = st.text_input(
            "Adresse",
            value=st.session_state.adresse,
            placeholder="Ville, pays",
        )

        st.session_state.linkedin = st.text_input(
            "LinkedIn",
            value=st.session_state.linkedin,
            placeholder="linkedin.com/in/...",
        )

        st.session_state.site = st.text_input(
            "Site web",
            value=st.session_state.site,
            placeholder="www.exemple.com",
        )

    with col2:

        st.markdown("### 📷 Photo")

        photo = st.file_uploader(
            "Ajouter une photo",
            type=["png", "jpg", "jpeg"],
            key="photo_uploader",
        )

        if photo is not None:
            st.session_state.photo = (
                photo.getvalue()
            )

        if st.session_state.photo:

            st.image(
                st.session_state.photo,
                width=180,
            )

            st.success(
                "Photo ajoutée."
            )


# =========================================================
# 2 - PROFIL
# =========================================================

with tabs[1]:

    section_card(
        "2",
        "Profil professionnel",
        COULEUR,
    )

    st.session_state.profil = st.text_area(
        "Présentez votre profil",
        value=st.session_state.profil,
        height=180,
        placeholder=(
            "Présentez votre formation, "
            "vos compétences, votre expérience "
            "et votre objectif professionnel."
        ),
    )


# =========================================================
# 3 - FORMATIONS
# =========================================================

with tabs[2]:

    section_card(
        "3",
        "Formations",
        COULEUR,
    )

    nombre_formations = st.number_input(
        "Nombre de formations",
        min_value=0,
        max_value=10,
        value=max(
            1,
            len(st.session_state.formation),
        ),
        step=1,
    )

    formations = []

    for i in range(
        int(nombre_formations)
    ):

        with st.expander(
            f"🎓 Formation {i + 1}",
            expanded=(i == 0),
        ):

            etablissement = st.text_input(
                "Établissement",
                key=f"formation_etablissement_{i}",
                placeholder="Université / École",
            )

            diplome = st.text_input(
                "Diplôme",
                key=f"formation_diplome_{i}",
                placeholder="Licence, Master, BTS...",
            )

            domaine = st.text_input(
                "Domaine / spécialité",
                key=f"formation_domaine_{i}",
                placeholder="Exemple : Études germaniques",
            )

            mois_debut, annee_debut = choisir_mois_annee(
                "Début",
                f"formation_debut_{i}",
            )

            en_cours = st.checkbox(
                "Formation en cours",
                key=f"formation_cours_{i}",
            )

            fin = ""

            if not en_cours:

                mois_fin, annee_fin = choisir_mois_annee(
                    "Fin",
                    f"formation_fin_{i}",
                )

                fin = format_mois_annee(
                    mois_fin,
                    annee_fin,
                )

            formations.append(
                {
                    "etablissement": etablissement,
                    "diplome": diplome,
                    "domaine": domaine,
                    "debut": format_mois_annee(
                        mois_debut,
                        annee_debut,
                    ),
                    "fin": (
                        "En cours"
                        if en_cours
                        else fin
                    ),
                }
            )

    st.session_state.formation = formations


# =========================================================
# 4 - EXPÉRIENCES
# =========================================================

with tabs[3]:

    section_card(
        "4",
        "Expériences professionnelles",
        COULEUR,
    )

    nombre_experiences = st.number_input(
        "Nombre d'expériences",
        min_value=0,
        max_value=10,
        value=len(
            st.session_state.experiences
        ),
        step=1,
    )

    experiences = []

    for i in range(
        int(nombre_experiences)
    ):

        with st.expander(
            f"💼 Expérience {i + 1}",
            expanded=(i == 0),
        ):

            poste = st.text_input(
                "Poste",
                key=f"experience_poste_{i}",
                placeholder="Intitulé du poste",
            )

            entreprise = st.text_input(
                "Entreprise / Organisation",
                key=f"experience_entreprise_{i}",
            )

            lieu = st.text_input(
                "Lieu",
                key=f"experience_lieu_{i}",
                placeholder="Ville, pays",
            )

            mois_debut, annee_debut = choisir_mois_annee(
                "Début",
                f"experience_debut_{i}",
            )

            actuel = st.checkbox(
                "Poste actuel",
                key=f"experience_actuel_{i}",
            )

            fin = ""

            if not actuel:

                mois_fin, annee_fin = choisir_mois_annee(
                    "Fin",
                    f"experience_fin_{i}",
                )

                fin = format_mois_annee(
                    mois_fin,
                    annee_fin,
                )

            description = st.text_area(
                "Description",
                key=f"experience_description_{i}",
                height=130,
                placeholder=(
                    "Décrivez vos principales "
                    "missions et réalisations."
                ),
            )

            experiences.append(
                {
                    "poste": poste,
                    "entreprise": entreprise,
                    "lieu": lieu,
                    "debut": format_mois_annee(
                        mois_debut,
                        annee_debut,
                    ),
                    "fin": (
                        "Poste actuel"
                        if actuel
                        else fin
                    ),
                    "description": description,
                }
            )

    st.session_state.experiences = experiences


# =========================================================
# 5 - COMPÉTENCES
# =========================================================

with tabs[4]:

    section_card(
        "5",
        "Compétences",
        COULEUR,
    )

    nombre_competences = st.number_input(
        "Nombre de compétences",
        min_value=0,
        max_value=20,
        value=len(
            st.session_state.competences
        ),
        step=1,
    )

    competences = []

    niveaux_competences = (
        [
            "Débutant",
            "Intermédiaire",
            "Avancé",
            "Expert",
        ]
        if st.session_state.langue_cv
        == "Français"
        else [
            "Beginner",
            "Intermediate",
            "Advanced",
            "Expert",
        ]
    )

    for i in range(
        int(nombre_competences)
    ):

        with st.expander(
            f"🛠️ Compétence {i + 1}",
            expanded=(i == 0),
        ):

            nom_competence = st.text_input(
                "Compétence",
                key=f"competence_nom_{i}",
                placeholder="Exemple : Analyse de données",
            )

            niveau = st.selectbox(
                "Niveau",
                niveaux_competences,
                key=f"competence_niveau_{i}",
            )

            competences.append(
                {
                    "nom": nom_competence,
                    "niveau": niveau,
                }
            )

    st.session_state.competences = competences


# =========================================================
# 6 - LANGUES
# =========================================================

with tabs[5]:

    section_card(
        "6",
        "Langues",
        COULEUR,
    )

    nombre_langues = st.number_input(
        "Nombre de langues",
        min_value=0,
        max_value=10,
        value=len(
            st.session_state.langues
        ),
        step=1,
    )

    langues = []

    for i in range(
        int(nombre_langues)
    ):

        with st.expander(
            f"🌍 Langue {i + 1}",
            expanded=(i == 0),
        ):

            nom_langue = st.text_input(
                "Langue",
                key=f"langue_nom_{i}",
                placeholder="Exemple : Allemand",
            )

            niveau_langue = st.selectbox(
                "Niveau",
                NIVEAUX[
                    st.session_state.langue_cv
                ],
                key=f"langue_niveau_{i}",
            )

            langues.append(
                {
                    "nom": nom_langue,
                    "niveau": niveau_langue,
                }
            )

    st.session_state.langues = langues


# =========================================================
# 7 - CERTIFICATIONS
# =========================================================

with tabs[6]:

    section_card(
        "7",
        "Certifications",
        COULEUR,
    )

    nombre_certifications = st.number_input(
        "Nombre de certifications",
        min_value=0,
        max_value=10,
        value=len(
            st.session_state.certifications
        ),
        step=1,
    )

    certifications = []

    for i in range(
        int(nombre_certifications)
    ):

        with st.expander(
            f"🏅 Certification {i + 1}",
            expanded=(i == 0),
        ):

            nom = st.text_input(
                "Nom de la certification",
                key=f"certification_nom_{i}",
            )

            organisme = st.text_input(
                "Organisme",
                key=f"certification_organisme_{i}",
            )

            date_obtention = st.text_input(
                "Date d'obtention",
                key=f"certification_date_{i}",
                placeholder="Exemple : Juin 2026",
            )

            certifications.append(
                {
                    "nom": nom,
                    "organisme": organisme,
                    "date": date_obtention,
                }
            )

    st.session_state.certifications = certifications


# =========================================================
# 8 - PROJETS
# =========================================================

with tabs[7]:

    section_card(
        "8",
        "Projets",
        COULEUR,
    )

    st.info(
        "Ajoutez ici vos projets académiques, "
        "professionnels, associatifs ou personnels."
    )

    nombre_projets = st.number_input(
        "Nombre de projets",
        min_value=0,
        max_value=10,
        value=len(
            st.session_state.projets
        ),
        step=1,
    )

    projets = []

    for i in range(
        int(nombre_projets)
    ):

        with st.expander(
            f"🚀 Projet {i + 1}",
            expanded=(i == 0),
        ):

            nom = st.text_input(
                "Nom du projet",
                key=f"projet_nom_{i}",
            )

            role = st.text_input(
                "Votre rôle",
                key=f"projet_role_{i}",
            )

            mois_debut, annee_debut = choisir_mois_annee(
                "Début",
                f"projet_debut_{i}",
            )

            actuel = st.checkbox(
                "Projet en cours",
                key=f"projet_actuel_{i}",
            )

            fin = ""

            if not actuel:

                mois_fin, annee_fin = choisir_mois_annee(
                    "Fin",
                    f"projet_fin_{i}",
                )

                fin = format_mois_annee(
                    mois_fin,
                    annee_fin,
                )

            description = st.text_area(
                "Description",
                key=f"projet_description_{i}",
                height=130,
            )

            projets.append(
                {
                    "nom": nom,
                    "role": role,
                    "debut": format_mois_annee(
                        mois_debut,
                        annee_debut,
                    ),
                    "fin": (
                        "En cours"
                        if actuel
                        else fin
                    ),
                    "description": description,
                }
            )

    st.session_state.projets = projets


# =========================================================
# 9 - ASSOCIATIONS
# =========================================================

with tabs[8]:

    section_card(
        "9",
        "Associations et engagements",
        COULEUR,
    )

    nombre_associations = st.number_input(
        "Nombre d'associations",
        min_value=0,
        max_value=10,
        value=len(
            st.session_state.associations
        ),
        step=1,
    )

    associations = []

    for i in range(
        int(nombre_associations)
    ):

        with st.expander(
            f"🤝 Association {i + 1}",
            expanded=(i == 0),
        ):

            organisation = st.text_input(
                "Organisation",
                key=f"association_organisation_{i}",
            )

            role = st.text_input(
                "Rôle",
                key=f"association_role_{i}",
            )

            description = st.text_area(
                "Description",
                key=f"association_description_{i}",
                height=120,
            )

            associations.append(
                {
                    "organisation": organisation,
                    "role": role,
                    "description": description,
                }
            )

    st.session_state.associations = associations


# =========================================================
# 10 - INTÉRÊTS
# =========================================================

with tabs[9]:

    section_card(
        "10",
        "Intérêts et loisirs",
        COULEUR,
    )

    nombre_interets = st.number_input(
        "Nombre d'intérêts",
        min_value=0,
        max_value=15,
        value=len(
            st.session_state.interets
        ),
        step=1,
    )

    interets = []

    for i in range(
        int(nombre_interets)
    ):

        interet = st.text_input(
            f"Intérêt {i + 1}",
            key=f"interet_{i}",
            placeholder="Exemple : Lecture",
        )

        interets.append(
            interet
        )

    st.session_state.interets = interets


# =========================================================
# 11 - ANALYSE D'OFFRE
# =========================================================

with tabs[10]:

    section_card(
        "11",
        "Analyse d'une offre d'emploi",
        COULEUR,
    )

    st.write(
        "Collez une offre d'emploi ci-dessous "
        "pour identifier les mots-clés les plus présents."
    )

    st.session_state.offre = st.text_area(
        "Offre d'emploi",
        value=st.session_state.offre,
        height=300,
        placeholder=(
            "Collez ici le texte de l'offre d'emploi..."
        ),
    )

    if st.button(
        "🔎 Analyser l'offre",
        use_container_width=True,
    ):

        texte = (
            st.session_state.offre
            .lower()
        )

        mots = re.findall(
            r"\b[a-zA-ZÀ-ÿ]{5,}\b",
            texte,
        )

        frequences = {}

        for mot in mots:

            frequences[mot] = (
                frequences.get(mot, 0)
                + 1
            )

        classement = sorted(
            frequences.items(),
            key=lambda x: x[1],
            reverse=True,
        )

        classement = [
            item
            for item in classement
            if item[1] >= 2
        ]

        if classement:

            st.success(
                f"{len(classement)} mots-clés "
                "identifiés."
            )

            for mot, nombre in classement[:30]:

                st.write(
                    f"**{mot}** : {nombre} occurrence(s)"
                )

        else:

            st.info(
                "Aucun mot-clé répété n'a été trouvé."
            )


# =========================================================
# CONSTRUCTION HTML DU CV
# =========================================================

def build_cv_html():

    couleur = st.session_state.couleur

    langue = st.session_state.langue_cv

    t = TEXTES[langue]

    nom = esc(
        st.session_state.nom
    )

    titre = esc(
        st.session_state.titre
    )

    # -----------------------------------------------------
    # CONTACT
    # -----------------------------------------------------

    contacts = []

    if st.session_state.email:
        contacts.append(
            f"✉ {esc(st.session_state.email)}"
        )

    if st.session_state.telephone:
        contacts.append(
            f"☎ {esc(st.session_state.telephone)}"
        )

    if st.session_state.adresse:
        contacts.append(
            f"⌖ {esc(st.session_state.adresse)}"
        )

    if st.session_state.linkedin:
        contacts.append(
            f"LinkedIn : {esc(st.session_state.linkedin)}"
        )

    if st.session_state.site:
        contacts.append(
            f"Web : {esc(st.session_state.site)}"
        )

    contact_html = "<br>".join(
        contacts
    )

    # -----------------------------------------------------
    # PHOTO
    # -----------------------------------------------------

    photo_html = ""

    if st.session_state.photo:

        photo_mime = "image/jpeg"

        if st.session_state.photo.startswith(
            b"\x89PNG"
        ):
            photo_mime = "image/png"

        elif st.session_state.photo.startswith(
            b"GIF"
        ):
            photo_mime = "image/gif"

        photo_base64 = base64.b64encode(
            st.session_state.photo
        ).decode()

        photo_html = f"""
        <img
            src="data:{photo_mime};base64,{photo_base64}"
            style="
                width:110px;
                height:110px;
                object-fit:cover;
                border-radius:50%;
                border:4px solid {couleur};
                float:right;
                margin-left:20px;
                margin-bottom:10px;
            "
        >
        """

    # -----------------------------------------------------
    # DÉBUT DU CV
    # -----------------------------------------------------

    html_cv = f"""
    <div class="cv-paper">

        {photo_html}

        <div
            class="cv-name"
            style="color:{couleur};"
        >
            {nom}
        </div>

        <div
            class="cv-title"
            style="color:#374151;"
        >
            {titre}
        </div>

        <div class="cv-contact">
            {contact_html}
        </div>

        <div
            class="cv-line"
            style="background:{couleur}; margin-top:15px;"
        >
        </div>
    """

    # -----------------------------------------------------
    # PROFIL
    # -----------------------------------------------------

    if st.session_state.profil:

        html_cv += f"""

        <div
            class="cv-section-title"
            style="color:{couleur};"
        >
            {t["profil"]}
        </div>

        <div class="cv-small">
            {esc(st.session_state.profil).replace(chr(10), "<br>")}
        </div>
        """

    # -----------------------------------------------------
    # FORMATIONS
    # -----------------------------------------------------

    if st.session_state.formation:

        html_cv += f"""

        <div
            class="cv-section-title"
            style="color:{couleur};"
        >
            {t["formation"]}
        </div>

        """

        for formation in st.session_state.formation:

            if not any(
                [
                    formation.get("etablissement"),
                    formation.get("diplome"),
                    formation.get("domaine"),
                ]
            ):
                continue

            html_cv += f"""

            <div class="cv-item">

                <div class="cv-date">
                    {esc(formation.get("debut", ""))}
                    –
                    {esc(formation.get("fin", ""))}
                </div>

                <strong>
                    {esc(formation.get("diplome", ""))}
                </strong>

                <div class="cv-small">
                    {esc(formation.get("domaine", ""))}
                </div>

                <div class="cv-small">
                    {esc(formation.get("etablissement", ""))}
                </div>

            </div>
            """

    # -----------------------------------------------------
    # EXPÉRIENCES
    # -----------------------------------------------------

    if st.session_state.experiences:

        html_cv += f"""

        <div
            class="cv-section-title"
            style="color:{couleur};"
        >
            {t["experience"]}
        </div>

        """

        for experience in st.session_state.experiences:

            if not any(
                [
                    experience.get("poste"),
                    experience.get("entreprise"),
                ]
            ):
                continue

            description = esc(
                experience.get(
                    "description",
                    "",
                )
            ).replace(
                "\n",
                "<br>",
            )

            html_cv += f"""

            <div class="cv-item">

                <div class="cv-date">
                    {esc(experience.get("debut", ""))}
                    –
                    {esc(experience.get("fin", ""))}
                </div>

                <strong>
                    {esc(experience.get("poste", ""))}
                </strong>

                <div class="cv-small">
                    {esc(experience.get("entreprise", ""))}
                    {(
                        " – "
                        + esc(experience.get("lieu", ""))
                    )
                    if experience.get("lieu")
                    else ""}
                </div>

                <div class="cv-small">
                    {description}
                </div>

            </div>
            """

    # -----------------------------------------------------
    # COMPÉTENCES
    # -----------------------------------------------------

    if st.session_state.competences:

        html_cv += f"""

        <div
            class="cv-section-title"
            style="color:{couleur};"
        >
            {t["competence"]}
        </div>

        """

        for competence in st.session_state.competences:

            if not competence.get("nom"):
                continue

            html_cv += f"""

            <span class="skill">
                {esc(competence.get("nom", ""))}
                –
                {esc(competence.get("niveau", ""))}
            </span>
            """

    # -----------------------------------------------------
    # LANGUES
    # -----------------------------------------------------

    if st.session_state.langues:

        html_cv += f"""

        <div
            class="cv-section-title"
            style="color:{couleur};"
        >
            {t["langue"]}
        </div>

        """

        for langue_item in st.session_state.langues:

            if not langue_item.get("nom"):
                continue

            html_cv += f"""

            <div class="cv-item">

                <strong>
                    {esc(langue_item.get("nom", ""))}
                </strong>

                <span class="cv-small">
                    – {esc(langue_item.get("niveau", ""))}
                </span>

            </div>
            """

    # -----------------------------------------------------
    # CERTIFICATIONS
    # -----------------------------------------------------

    if st.session_state.certifications:

        html_cv += f"""

        <div
            class="cv-section-title"
            style="color:{couleur};"
        >
            {t["certification"]}
        </div>

        """

        for certification in st.session_state.certifications:

            if not certification.get("nom"):
                continue

            html_cv += f"""

            <div class="cv-item">

                <strong>
                    {esc(certification.get("nom", ""))}
                </strong>

                <div class="cv-small">
                    {esc(certification.get("organisme", ""))}
                    {(
                        " – "
                        + esc(certification.get("date", ""))
                    )
                    if certification.get("date")
                    else ""}
                </div>

            </div>
            """

    # -----------------------------------------------------
    # PROJETS
    # -----------------------------------------------------

    if st.session_state.projets:

        html_cv += f"""

        <div
            class="cv-section-title"
            style="color:{couleur};"
        >
            {t["projet"]}
        </div>

        """

        for projet in st.session_state.projets:

            if not projet.get("nom"):
                continue

            description = esc(
                projet.get(
                    "description",
                    "",
                )
            ).replace(
                "\n",
                "<br>",
            )

            html_cv += f"""

            <div class="cv-item">

                <div class="cv-date">
                    {esc(projet.get("debut", ""))}
                    –
                    {esc(projet.get("fin", ""))}
                </div>

                <strong>
                    {esc(projet.get("nom", ""))}
                </strong>

                <div class="cv-small">
                    {esc(projet.get("role", ""))}
                </div>

                <div class="cv-small">
                    {description}
                </div>

            </div>
            """

    # -----------------------------------------------------
    # ASSOCIATIONS
    # -----------------------------------------------------

    if st.session_state.associations:

        html_cv += f"""

        <div
            class="cv-section-title"
            style="color:{couleur};"
        >
            {t["association"]}
        </div>

        """

        for association in st.session_state.associations:

            if not association.get(
                "organisation"
            ):
                continue

            description = esc(
                association.get(
                    "description",
                    "",
                )
            ).replace(
                "\n",
                "<br>",
            )

            html_cv += f"""

            <div class="cv-item">

                <strong>
                    {esc(association.get("organisation", ""))}
                </strong>

                <div class="cv-small">
                    {esc(association.get("role", ""))}
                </div>

                <div class="cv-small">
                    {description}
                </div>

            </div>
            """

    # -----------------------------------------------------
    # INTÉRÊTS
    # -----------------------------------------------------

    if st.session_state.interets:

        interets_valides = [
            esc(i)
            for i in st.session_state.interets
            if i
        ]

        if interets_valides:

            html_cv += f"""

            <div
                class="cv-section-title"
                style="color:{couleur};"
            >
                {t["interet"]}
            </div>

            <div class="cv-small">
                {" • ".join(interets_valides)}
            </div>
            """

    # -----------------------------------------------------
    # FIN
    # -----------------------------------------------------

    html_cv += """

    </div>
    """

    return html_cv


# =========================================================
# EXPORT WORD
# =========================================================

def generate_word():

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    couleur = couleur_rgb(
        st.session_state.couleur
    )

    # -----------------------------------------------------
    # NOM
    # -----------------------------------------------------

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        st.session_state.nom
    )

    run.bold = True
    run.font.size = Pt(25)
    run.font.color.rgb = couleur

    # -----------------------------------------------------
    # TITRE
    # -----------------------------------------------------

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        st.session_state.titre
    )

    run.bold = True
    run.font.size = Pt(14)

    # -----------------------------------------------------
    # CONTACT
    # -----------------------------------------------------

    contacts = []

    if st.session_state.email:
        contacts.append(
            st.session_state.email
        )

    if st.session_state.telephone:
        contacts.append(
            st.session_state.telephone
        )

    if st.session_state.adresse:
        contacts.append(
            st.session_state.adresse
        )

    if st.session_state.linkedin:
        contacts.append(
            st.session_state.linkedin
        )

    if st.session_state.site:
        contacts.append(
            st.session_state.site
        )

    if contacts:

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            " | ".join(contacts)
        )

        run.font.size = Pt(9)

    # -----------------------------------------------------
    # FONCTION SECTION
    # -----------------------------------------------------

    def word_section(titre):

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(3)

        run = paragraph.add_run(
            titre.upper()
        )

        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = couleur

        try:

            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn

            p = paragraph._p

            pPr = p.get_or_add_pPr()

            pBdr = OxmlElement(
                "w:pBdr"
            )

            bottom = OxmlElement(
                "w:bottom"
            )

            bottom.set(
                qn("w:val"),
                "single",
            )

            bottom.set(
                qn("w:sz"),
                "6",
            )

            bottom.set(
                qn("w:space"),
                "1",
            )

            bottom.set(
                qn("w:color"),
                st.session_state.couleur.replace(
                    "#",
                    "",
                ),
            )

            pBdr.append(bottom)

            pPr.append(pBdr)

        except Exception:
            pass

    # -----------------------------------------------------
    # PROFIL
    # -----------------------------------------------------

    if st.session_state.profil:

        word_section(
            TEXTES[
                st.session_state.langue_cv
            ]["profil"]
        )

        p = document.add_paragraph(
            st.session_state.profil
        )

        for run in p.runs:
            run.font.size = Pt(9)

    # -----------------------------------------------------
    # FORMATIONS
    # -----------------------------------------------------

    if st.session_state.formation:

        word_section(
            TEXTES[
                st.session_state.langue_cv
            ]["formation"]
        )

        for formation in st.session_state.formation:

            p = document.add_paragraph()

            run = p.add_run(
                formation.get(
                    "diplome",
                    "",
                )
            )

            run.bold = True
            run.font.size = Pt(10)

            p.add_run(
                f" – {formation.get('domaine', '')}"
            )

            p.add_run(
                f"\n{formation.get('etablissement', '')}"
            )

            p.add_run(
                f"\n{formation.get('debut', '')} – "
                f"{formation.get('fin', '')}"
            )

    # -----------------------------------------------------
    # EXPÉRIENCES
    # -----------------------------------------------------

    if st.session_state.experiences:

        word_section(
            TEXTES[
                st.session_state.langue_cv
            ]["experience"]
        )

        for experience in st.session_state.experiences:

            p = document.add_paragraph()

            run = p.add_run(
                experience.get(
                    "poste",
                    "",
                )
            )

            run.bold = True

            p.add_run(
                f" – {experience.get('entreprise', '')}"
            )

            if experience.get("lieu"):
                p.add_run(
                    f" – {experience.get('lieu')}"
                )

            p.add_run(
                f"\n{experience.get('debut', '')} – "
                f"{experience.get('fin', '')}"
            )

            if experience.get(
                "description"
            ):
                p.add_run(
                    "\n"
                    + experience.get(
                        "description",
                        "",
                    )
                )

    # -----------------------------------------------------
    # COMPÉTENCES
    # -----------------------------------------------------

    if st.session_state.competences:

        word_section(
            TEXTES[
                st.session_state.langue_cv
            ]["competence"]
        )

        for competence in st.session_state.competences:

            p = document.add_paragraph(
                style=None
            )

            p.add_run(
                f"• {competence.get('nom', '')} "
            )

            p.add_run(
                f"({competence.get('niveau', '')})"
            )

    # -----------------------------------------------------
    # LANGUES
    # -----------------------------------------------------

    if st.session_state.langues:

        word_section(
            TEXTES[
                st.session_state.langue_cv
            ]["langue"]
        )

        for langue in st.session_state.langues:

            p = document.add_paragraph()

            p.add_run(
                f"• {langue.get('nom', '')} "
            )

            p.add_run(
                f"({langue.get('niveau', '')})"
            )

    # -----------------------------------------------------
    # CERTIFICATIONS
    # -----------------------------------------------------

    if st.session_state.certifications:

        word_section(
            TEXTES[
                st.session_state.langue_cv
            ]["certification"]
        )

        for certification in st.session_state.certifications:

            p = document.add_paragraph()

            p.add_run(
                certification.get(
                    "nom",
                    "",
                )
            ).bold = True

            p.add_run(
                f" – {certification.get('organisme', '')}"
            )

            if certification.get("date"):
                p.add_run(
                    f" – {certification.get('date')}"
                )

    # -----------------------------------------------------
    # PROJETS
    # -----------------------------------------------------

    if st.session_state.projets:

        word_section(
            TEXTES[
                st.session_state.langue_cv
            ]["projet"]
        )

        for projet in st.session_state.projets:

            p = document.add_paragraph()

            p.add_run(
                projet.get(
                    "nom",
                    "",
                )
            ).bold = True

            if projet.get("role"):
                p.add_run(
                    f" – {projet.get('role')}"
                )

            p.add_run(
                f"\n{projet.get('debut', '')} – "
                f"{projet.get('fin', '')}"
            )

            if projet.get("description"):
                p.add_run(
                    "\n"
                    + projet.get(
                        "description",
                        "",
                    )
                )

    # -----------------------------------------------------
    # ASSOCIATIONS
    # -----------------------------------------------------

    if st.session_state.associations:

        word_section(
            TEXTES[
                st.session_state.langue_cv
            ]["association"]
        )

        for association in st.session_state.associations:

            p = document.add_paragraph()

            p.add_run(
                association.get(
                    "organisation",
                    "",
                )
            ).bold = True

            if association.get("role"):
                p.add_run(
                    f" – {association.get('role')}"
                )

            if association.get(
                "description"
            ):
                p.add_run(
                    "\n"
                    + association.get(
                        "description",
                        "",
                    )
                )

    # -----------------------------------------------------
    # INTÉRÊTS
    # -----------------------------------------------------

    interets = [
        i
        for i in st.session_state.interets
        if i
    ]

    if interets:

        word_section(
            TEXTES[
                st.session_state.langue_cv
            ]["interet"]
        )

        document.add_paragraph(
            " • ".join(interets)
        )

    # -----------------------------------------------------
    # SAUVEGARDE
    # -----------------------------------------------------

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output


# =========================================================
# EXPORT PDF
# =========================================================

def generate_pdf():

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=35,
        leftMargin=35,
        topMargin=35,
        bottomMargin=35,
    )

    styles = getSampleStyleSheet()

    couleur = HexColor(
        st.session_state.couleur
    )

    name_style = ParagraphStyle(
        "CVName",
        parent=styles["Normal"],
        fontSize=24,
        leading=28,
        alignment=1,
        textColor=couleur,
        spaceAfter=4,
    )

    title_style = ParagraphStyle(
        "CVTitle",
        parent=styles["Normal"],
        fontSize=11,
        leading=14,
        alignment=1,
        spaceAfter=7,
    )

    contact_style = ParagraphStyle(
        "CVContact",
        parent=styles["Normal"],
        fontSize=8,
        leading=11,
        alignment=1,
        textColor=colors.HexColor(
            "#374151"
        ),
        spaceAfter=8,
    )

    section_style = ParagraphStyle(
        "CVSection",
        parent=styles["Normal"],
        fontSize=11,
        leading=14,
        textColor=couleur,
        spaceBefore=8,
        spaceAfter=4,
    )

    body_style = ParagraphStyle(
        "CVBody",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        spaceAfter=4,
    )

    small_style = ParagraphStyle(
        "CVSmall",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor(
            "#4b5563"
        ),
        spaceAfter=3,
    )

    story = []

    # -----------------------------------------------------
    # NOM
    # -----------------------------------------------------

    story.append(
        Paragraph(
            esc(
                st.session_state.nom
            ),
            name_style,
        )
    )

    story.append(
        Paragraph(
            esc(
                st.session_state.titre
            ),
            title_style,
        )
    )

    # -----------------------------------------------------
    # CONTACT
    # -----------------------------------------------------

    contacts = []

    if st.session_state.email:
        contacts.append(
            esc(
                st.session_state.email
            )
        )

    if st.session_state.telephone:
        contacts.append(
            esc(
                st.session_state.telephone
            )
        )

    if st.session_state.adresse:
        contacts.append(
            esc(
                st.session_state.adresse
            )
        )

    if st.session_state.linkedin:
        contacts.append(
            esc(
                st.session_state.linkedin
            )
        )

    if st.session_state.site:
        contacts.append(
            esc(
                st.session_state.site
            )
        )

    if contacts:

        story.append(
            Paragraph(
                " | ".join(contacts),
                contact_style,
            )
        )

    story.append(
        HRFlowable(
            width="100%",
            thickness=1.5,
            color=couleur,
            spaceAfter=8,
        )
    )

    # -----------------------------------------------------
    # PROFIL
    # -----------------------------------------------------

    if st.session_state.profil:

        story.append(
            Paragraph(
                esc(
                    TEXTES[
                        st.session_state.langue_cv
                    ]["profil"]
                ),
                section_style,
            )
        )

        story.append(
            Paragraph(
                esc(
                    st.session_state.profil
                ).replace(
                    "\n",
                    "<br/>",
                ),
                body_style,
            )
        )

    # -----------------------------------------------------
    # FORMATIONS
    # -----------------------------------------------------

    if st.session_state.formation:

        story.append(
            Paragraph(
                esc(
                    TEXTES[
                        st.session_state.langue_cv
                    ]["formation"]
                ),
                section_style,
            )
        )

        for formation in st.session_state.formation:

            if not formation.get(
                "diplome"
            ):
                continue

            story.append(
                Paragraph(
                    f"<b>{esc(formation.get('diplome', ''))}</b>",
                    body_style,
                )
            )

            story.append(
                Paragraph(
                    esc(
                        formation.get(
                            "domaine",
                            "",
                        )
                    ),
                    small_style,
                )
            )

            story.append(
                Paragraph(
                    esc(
                        formation.get(
                            "etablissement",
                            "",
                        )
                    ),
                    small_style,
                )
            )

            story.append(
                Paragraph(
                    f"{esc(formation.get('debut', ''))} – "
                    f"{esc(formation.get('fin', ''))}",
                    small_style,
                )
            )

    # -----------------------------------------------------
    # EXPÉRIENCES
    # -----------------------------------------------------

    if st.session_state.experiences:

        story.append(
            Paragraph(
                esc(
                    TEXTES[
                        st.session_state.langue_cv
                    ]["experience"]
                ),
                section_style,
            )
        )

        for experience in st.session_state.experiences:

            if not experience.get(
                "poste"
            ):
                continue

            story.append(
                Paragraph(
                    f"<b>{esc(experience.get('poste', ''))}</b> "
                    f"– {esc(experience.get('entreprise', ''))}",
                    body_style,
                )
            )

            if experience.get("lieu"):

                story.append(
                    Paragraph(
                        esc(
                            experience.get(
                                "lieu",
                                "",
                            )
                        ),
                        small_style,
                    )
                )

            story.append(
                Paragraph(
                    f"{esc(experience.get('debut', ''))} – "
                    f"{esc(experience.get('fin', ''))}",
                    small_style,
                )
            )

            if experience.get(
                "description"
            ):

                story.append(
                    Paragraph(
                        esc(
                            experience.get(
                                "description",
                                "",
                            )
                        ).replace(
                            "\n",
                            "<br/>",
                        ),
                        small_style,
                    )
                )

    # -----------------------------------------------------
    # COMPÉTENCES
    # -----------------------------------------------------

    if st.session_state.competences:

        story.append(
            Paragraph(
                esc(
                    TEXTES[
                        st.session_state.langue_cv
                    ]["competence"]
                ),
                section_style,
            )
        )

        skills = []

        for competence in st.session_state.competences:

            if competence.get("nom"):

                skills.append(
                    f"<b>{esc(competence.get('nom', ''))}</b> "
                    f"({esc(competence.get('niveau', ''))})"
                )

        if skills:

            story.append(
                Paragraph(
                    " • ".join(skills),
                    body_style,
                )
            )

    # -----------------------------------------------------
    # LANGUES
    # -----------------------------------------------------

    if st.session_state.langues:

        story.append(
            Paragraph(
                esc(
                    TEXTES[
                        st.session_state.langue_cv
                    ]["langue"]
                ),
                section_style,
            )
        )

        for langue in st.session_state.langues:

            if not langue.get("nom"):
                continue

            story.append(
                Paragraph(
                    f"<b>{esc(langue.get('nom', ''))}</b> "
                    f"– {esc(langue.get('niveau', ''))}",
                    body_style,
                )
            )

    # -----------------------------------------------------
    # CERTIFICATIONS
    # -----------------------------------------------------

    if st.session_state.certifications:

        story.append(
            Paragraph(
                esc(
                    TEXTES[
                        st.session_state.langue_cv
                    ]["certification"]
                ),
                section_style,
            )
        )

        for certification in st.session_state.certifications:

            if not certification.get(
                "nom"
            ):
                continue

            texte = (
                f"<b>{esc(certification.get('nom', ''))}</b>"
            )

            if certification.get(
                "organisme"
            ):

                texte += (
                    f" – "
                    f"{esc(certification.get('organisme'))}"
                )

            if certification.get(
                "date"
            ):

                texte += (
                    f" – "
                    f"{esc(certification.get('date'))}"
                )

            story.append(
                Paragraph(
                    texte,
                    body_style,
                )
            )

    # -----------------------------------------------------
    # PROJETS
    # -----------------------------------------------------

    if st.session_state.projets:

        story.append(
            Paragraph(
                esc(
                    TEXTES[
                        st.session_state.langue_cv
                    ]["projet"]
                ),
                section_style,
            )
        )

        for projet in st.session_state.projets:

            if not projet.get(
                "nom"
            ):
                continue

            story.append(
                Paragraph(
                    f"<b>{esc(projet.get('nom', ''))}</b> "
                    f"– {esc(projet.get('role', ''))}",
                    body_style,
                )
            )

            story.append(
                Paragraph(
                    f"{esc(projet.get('debut', ''))} – "
                    f"{esc(projet.get('fin', ''))}",
                    small_style,
                )
            )

            if projet.get(
                "description"
            ):

                story.append(
                    Paragraph(
                        esc(
                            projet.get(
                                "description",
                                "",
                            )
                        ).replace(
                            "\n",
                            "<br/>",
                        ),
                        small_style,
                    )
                )

    # -----------------------------------------------------
    # ASSOCIATIONS
    # -----------------------------------------------------

    if st.session_state.associations:

        story.append(
            Paragraph(
                esc(
                    TEXTES[
                        st.session_state.langue_cv
                    ]["association"]
                ),
                section_style,
            )
        )

        for association in st.session_state.associations:

            if not association.get(
                "organisation"
            ):
                continue

            story.append(
                Paragraph(
                    f"<b>{esc(association.get('organisation', ''))}</b> "
                    f"– {esc(association.get('role', ''))}",
                    body_style,
                )
            )

            if association.get(
                "description"
            ):

                story.append(
                    Paragraph(
                        esc(
                            association.get(
                                "description",
                                "",
                            )
                        ).replace(
                            "\n",
                            "<br/>",
                        ),
                        small_style,
                    )
                )

    # -----------------------------------------------------
    # INTÉRÊTS
    # -----------------------------------------------------

    interets = [
        esc(i)
        for i in st.session_state.interets
        if i
    ]

    if interets:

        story.append(
            Paragraph(
                esc(
                    TEXTES[
                        st.session_state.langue_cv
                    ]["interet"]
                ),
                section_style,
            )
        )

        story.append(
            Paragraph(
                " • ".join(interets),
                body_style,
            )
        )

    # -----------------------------------------------------
    # GÉNÉRATION
    # -----------------------------------------------------

    document.build(story)

    output.seek(0)

    return output


# =========================================================
# SCORE ATS
# =========================================================

def calcul_ats():

    score = 0

    details = []

    if st.session_state.nom:

        score += 10
        details.append(
            "Nom complet"
        )

    if st.session_state.titre:

        score += 10
        details.append(
            "Titre professionnel"
        )

    if st.session_state.email:

        score += 10
        details.append(
            "Adresse e-mail"
        )

    if st.session_state.telephone:

        score += 10
        details.append(
            "Numéro de téléphone"
        )

    if st.session_state.profil:

        score += 10
        details.append(
            "Profil professionnel"
        )

    if any(
        formation.get("diplome")
        for formation
        in st.session_state.formation
    ):

        score += 10
        details.append(
            "Formation"
        )

    if any(
        experience.get("poste")
        for experience
        in st.session_state.experiences
    ):

        score += 10
        details.append(
            "Expérience professionnelle"
        )

    if any(
        competence.get("nom")
        for competence
        in st.session_state.competences
    ):

        score += 10
        details.append(
            "Compétences"
        )

    if any(
        langue.get("nom")
        for langue
        in st.session_state.langues
    ):

        score += 10
        details.append(
            "Langues"
        )

    if (
        st.session_state.certifications
        or st.session_state.projets
    ):

        score += 10
        details.append(
            "Certifications ou projets"
        )

    return score, details


# =========================================================
# 12 - APERÇU ET EXPORT
# =========================================================

with tabs[11]:

    section_card(
        "12",
        "Aperçu & Export",
        COULEUR,
    )

    col1, col2 = st.columns(
        [3, 1]
    )

    with col1:

        st.markdown(
            "### 👀 Aperçu du CV"
        )

    with col2:

        score, details = calcul_ats()

        st.metric(
            "Score CV",
            f"{score}/100",
        )

    # -----------------------------------------------------
    # APERÇU
    # -----------------------------------------------------

    st.html(
        build_cv_html()
    )

    st.markdown("---")

    st.markdown(
        "### 📥 Exporter votre CV"
    )

    col1, col2 = st.columns(2)

    with col1:

        word_file = generate_word()

        st.download_button(
            "📄 Télécharger en Word",
            data=word_file.getvalue(),
            file_name="CV_Professionnel.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            use_container_width=True,
        )

    with col2:

        pdf_file = generate_pdf()

        st.download_button(
            "📑 Télécharger en PDF",
            data=pdf_file.getvalue(),
            file_name="CV_Professionnel.pdf",
            mime="application/pdf",
            use_container_width=True,
        )

    # -----------------------------------------------------
    # ANALYSE
    # -----------------------------------------------------

    st.markdown("---")

    st.markdown(
        "### 📊 Analyse rapide du CV"
    )

    if score >= 80:

        st.success(
            "Excellent ! Votre CV contient "
            "la plupart des informations importantes."
        )

    elif score >= 60:

        st.warning(
            "Votre CV est bien avancé. "
            "Ajoutez quelques informations "
            "pour améliorer son contenu."
        )

    else:

        st.info(
            "Votre CV peut encore être complété. "
            "Ajoutez vos informations principales."
        )

    if details:

        st.write(
            "Éléments détectés :"
        )

        for detail in details:

            st.write(
                f"✓ {detail}"
            )